"""
Прототип парсера каталога Дикси (Москва, дефолтный регион).

Этапы:
1. Sitemap https://dixy.ru/sitemap-iblock-1.xml -> список product URL.
2. Карточка /product/<slug>-<sku>/ -> module.initElementData({...}) + DOM (характеристики, состав, срок хранения).
3. /ajax/reviews-json.php -> агрегат (распределение по звёздам) + 10 свежих отзывов.
4. Инкрементная запись в products.jsonl + reviews.jsonl (поддерживает перезапуск).
5. Экспорт в dixy_catalog.xlsx (два листа: products / reviews).

Запуск:
    python parser.py --limit 50          # тестовый прогон
    python parser.py --limit 0           # без лимита (~30k SKU)
    python parser.py --export-only       # только пересобрать xlsx из jsonl
"""

import argparse
import json
import random
import re
import threading
import time
from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import datetime, timezone
from pathlib import Path
from xml.etree import ElementTree as ET

from curl_cffi import requests
from selectolax.parser import HTMLParser

HERE = Path(__file__).resolve().parent
BASE = "https://dixy.ru"
SITEMAP_URL = f"{BASE}/sitemap-iblock-1.xml"
PRODUCTS_JSONL = HERE / "products.jsonl"
REVIEWS_JSONL = HERE / "reviews.jsonl"
XLSX_PATH = HERE / "dixy_catalog.xlsx"

PRODUCT_COLUMNS = [
    "sku", "internal_id", "slug", "url", "name", "brand", "type",
    "section", "section_id", "parent_section", "weight",
    "price", "old_price", "has_discount", "discount_pct",
    "badge_title", "badge_class", "price_block_class",
    "rating", "reviews_total",
    "reviews_5", "reviews_4", "reviews_3", "reviews_2", "reviews_1",
    "calories", "protein", "fat", "carbs",
    "composition",
    "shelf_life", "storage_temp_min", "storage_temp_max",
    "country", "packaging_type", "packaging_form", "cooking_time", "trademark",
    "stock_quantity", "status",
    "is_alco", "is_energy_drink", "is_gift",
    "preview_image", "detail_image", "all_images",
    "description",
    "parsed_at",
]

REVIEW_COLUMNS = ["sku", "review_id", "date", "user_name", "rating", "comment"]


# ---------- HTTP ----------

def _load_cookies_into(s):
    cookies_file = HERE / "cookies.txt"
    if not cookies_file.exists():
        return
    for line in cookies_file.read_text(encoding="utf-8").splitlines():
        if not line or line.startswith("#") and not line.startswith("#HttpOnly_"):
            continue
        parts = line.lstrip("#HttpOnly_").split("\t")
        if len(parts) == 7:
            _domain, _flag, _path, _secure, _exp, name, value = parts
            s.cookies.set(name, value, domain=".dixy.ru")


def make_session(proxy: str | None = None):
    kwargs = {"impersonate": "chrome"}
    if proxy:
        kwargs["proxies"] = {"http": proxy, "https": proxy}
    s = requests.Session(**kwargs)
    s.headers.update({"Accept-Language": "ru-RU,ru;q=0.9,en;q=0.8"})
    _load_cookies_into(s)
    return s




def polite_sleep(min_s=0.6, max_s=1.4):
    time.sleep(random.uniform(min_s, max_s))


# ---------- Sitemap ----------

def fetch_sitemap_urls(session, limit=None):
    local = HERE / "sitemap-iblock-1.xml"
    if local.exists():
        print(f"[sitemap] using local {local.name} ({local.stat().st_size // 1024} KB)", flush=True)
        data = local.read_bytes()
    else:
        r = session.get(SITEMAP_URL, timeout=60)
        r.raise_for_status()
        data = r.content
    ns = {"sm": "http://www.sitemaps.org/schemas/sitemap/0.9"}
    root = ET.fromstring(data)
    urls = [u.text.strip() for u in root.findall("sm:url/sm:loc", ns) if u.text]
    urls = [u for u in urls if "/product/" in u]
    if limit and limit > 0:
        urls = urls[:limit]
    return urls


# ---------- Product page ----------

_RE_ELEM_DATA = re.compile(r"module\.initElementData\(")


def _extract_init_element_data(html: str) -> dict | None:
    m = _RE_ELEM_DATA.search(html)
    if not m:
        return None
    start = m.end()
    decoder = json.JSONDecoder()
    try:
        obj, _ = decoder.raw_decode(html, start)
        return obj
    except json.JSONDecodeError:
        return None


def _text_or_none(node):
    return node.text(strip=True) if node else None


def _parse_rendered_props(rendered_html: str) -> dict:
    out = {
        "composition": None,
        "shelf_life": None,
        "storage_temp_min": None,
        "storage_temp_max": None,
        "country": None,
        "packaging_type": None,
        "packaging_form": None,
        "cooking_time": None,
        "trademark": None,
        "type_from_props": None,
    }
    if not rendered_html:
        return out
    tree = HTMLParser(rendered_html)
    for h in tree.css(".detail-data__holder"):
        title_el = h.css_first(".title")
        if not title_el:
            continue
        title = title_el.text(strip=True).lower()
        if "состав" in title:
            text_el = h.css_first(".text")
            if text_el:
                out["composition"] = text_el.text(strip=True)
        elif "условия хранения" in title:
            for line in h.css(".list__line"):
                lbl_el = line.css_first(".text")
                val_el = line.css_first(".list__link span") or line.css_first("span")
                if not lbl_el or not val_el:
                    continue
                lbl = lbl_el.text(strip=True).lower()
                val = val_el.text(strip=True)
                if "срок хранения" in lbl:
                    out["shelf_life"] = val
                elif "температура хранения мин" in lbl:
                    out["storage_temp_min"] = val
                elif "температура хранения макс" in lbl:
                    out["storage_temp_max"] = val
        elif "информация" in title:
            for line in h.css(".list__line"):
                spans = line.css("span")
                if len(spans) < 2:
                    continue
                lbl = spans[0].text(strip=True).lower()
                val = spans[-1].text(strip=True)
                if "торговая марка" in lbl:
                    out["trademark"] = val
                elif "страна" in lbl:
                    out["country"] = val
                elif "тип товара" in lbl:
                    out["type_from_props"] = val
                elif "время приготовления" in lbl:
                    out["cooking_time"] = val
                elif "фасовка" in lbl:
                    out["packaging_form"] = val
                elif "тип упаковки" in lbl:
                    out["packaging_type"] = val
    return out


def parse_product(html: str, url: str) -> dict | None:
    data = _extract_init_element_data(html)
    if not data:
        return None

    def _as_list_of_dicts(v):
        return [x for x in v if isinstance(x, dict)] if isinstance(v, list) else []

    food = {f["name"].lower(): f.get("value") for f in _as_list_of_dicts(data.get("food_price")) if "name" in f}
    badges = _as_list_of_dicts(data.get("badges"))
    first_badge = badges[0] if badges else {}

    previews = [p["src"] for p in _as_list_of_dicts(data.get("previews")) if p.get("src")]
    details = [p["src"] for p in _as_list_of_dicts(data.get("details")) if p.get("src")]
    all_images = sorted({*previews, *details})

    rendered_props = data.get("renderedProps") or ""
    props = _parse_rendered_props(rendered_props)

    discount_pct = None
    if first_badge.get("num"):
        m = re.search(r"-?(\d+)", first_badge["num"])
        if m:
            discount_pct = int(m.group(1))

    sku = data.get("xml_id")
    slug = data.get("code")

    return {
        "sku": sku,
        "internal_id": data.get("id"),
        "slug": slug,
        "url": url,
        "name": data.get("title"),
        "brand": data.get("brand"),
        "type": data.get("type") or props.get("type_from_props"),
        "section": data.get("section"),
        "section_id": data.get("section_id"),
        "parent_section": data.get("parent_section") or None,
        "weight": data.get("weight"),
        "price": _to_float(data.get("priceSimple")),
        "old_price": _to_float(data.get("oldPriceSimple")),
        "has_discount": bool(data.get("crossPrice")),
        "discount_pct": discount_pct,
        "badge_title": first_badge.get("title"),
        "badge_class": first_badge.get("class"),
        "price_block_class": data.get("status"),
        "rating": _to_float(data.get("rating")),
        # reviews_* и reviews_total заполняются позже из AJAX
        "calories": food.get("калории"),
        "protein": food.get("белки"),
        "fat": food.get("жиры"),
        "carbs": food.get("углеводы"),
        "composition": props.get("composition"),
        "shelf_life": props.get("shelf_life"),
        "storage_temp_min": props.get("storage_temp_min"),
        "storage_temp_max": props.get("storage_temp_max"),
        "country": props.get("country"),
        "packaging_type": props.get("packaging_type"),
        "packaging_form": props.get("packaging_form"),
        "cooking_time": props.get("cooking_time"),
        "trademark": props.get("trademark"),
        "stock_quantity": data.get("stock_quantity"),
        "status": data.get("status"),
        "is_alco": bool(data.get("isAlko")),
        "is_energy_drink": bool(data.get("isEnergyDrink")),
        "is_gift": bool(data.get("isGift")),
        "preview_image": _abs_url(previews[0] if previews else None),
        "detail_image": _abs_url(data.get("detail_src")),
        "all_images": ";".join(_abs_url(u) for u in all_images),
        "description": data.get("text"),
        "parsed_at": datetime.now(timezone.utc).isoformat(timespec="seconds"),
    }


def _to_float(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", "."))
    except ValueError:
        return None


def _abs_url(u):
    if not u:
        return None
    if u.startswith("http"):
        return u
    return BASE + u


# ---------- Reviews ----------

def fetch_reviews(session, xml_id: str):
    url = f"{BASE}/ajax/reviews-json.php"
    params = {"sort": "date", "order": "desc", "page_size": 10, "page": 1, "xml_id": xml_id}
    r = session.get(url, params=params, timeout=30)
    if r.status_code != 200:
        return None
    try:
        data = r.json()
    except Exception:
        return None
    return data


# ---------- IO ----------

def load_seen_skus(jsonl_path: Path) -> set[str]:
    if not jsonl_path.exists():
        return set()
    seen = set()
    with jsonl_path.open(encoding="utf-8") as f:
        for line in f:
            try:
                obj = json.loads(line)
                if obj.get("sku"):
                    seen.add(str(obj["sku"]))
            except json.JSONDecodeError:
                continue
    return seen


def append_jsonl(path: Path, obj):
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(obj, ensure_ascii=False) + "\n")


# ---------- Excel ----------

def export_xlsx():
    from openpyxl import Workbook

    wb = Workbook()
    ws_p = wb.active
    ws_p.title = "products"
    ws_p.append(PRODUCT_COLUMNS)

    if PRODUCTS_JSONL.exists():
        with PRODUCTS_JSONL.open(encoding="utf-8") as f:
            for line in f:
                obj = json.loads(line)
                ws_p.append([obj.get(c) for c in PRODUCT_COLUMNS])

    ws_r = wb.create_sheet("reviews")
    ws_r.append(REVIEW_COLUMNS)
    if REVIEWS_JSONL.exists():
        with REVIEWS_JSONL.open(encoding="utf-8") as f:
            for line in f:
                obj = json.loads(line)
                ws_r.append([obj.get(c) for c in REVIEW_COLUMNS])

    for ws in (ws_p, ws_r):
        for col_idx, name in enumerate(ws[1], start=1):
            ws.column_dimensions[name.column_letter].width = max(12, min(40, len(str(name.value)) + 2))

    wb.save(XLSX_PATH)
    print(f"[xlsx] {XLSX_PATH} - products={ws_p.max_row - 1}, reviews={ws_r.max_row - 1}")


# ---------- Main ----------

_thread_local = threading.local()


def _get_thread_session(proxy):
    s = getattr(_thread_local, "session", None)
    if s is None:
        s = make_session(proxy=proxy)
        _thread_local.session = s
    return s


def _process_one(url, idx, total, proxy, counters, lock):
    try:
        session = _get_thread_session(proxy)
        r = session.get(url, timeout=30)
        if r.status_code != 200:
            print(f"[{idx}/{total}] HTTP {r.status_code} {url}", flush=True)
            with lock:
                counters["failed"] += 1
            return
        product = parse_product(r.text, url)
        if not product or not product.get("sku"):
            print(f"[{idx}/{total}] no initElementData on {url}", flush=True)
            with lock:
                counters["failed"] += 1
            return
        sku = str(product["sku"])

        review_rows = []
        try:
            rr = session.get(
                f"{BASE}/ajax/reviews-json.php",
                params={"sort": "date", "order": "desc", "page_size": 10, "page": 1, "xml_id": sku},
                timeout=30,
            )
            if rr.status_code == 200:
                payload = rr.json()
                q = payload.get("quantity") or {}
                product["reviews_total"] = q.get("all")
                product["reviews_5"] = q.get("5")
                product["reviews_4"] = q.get("4")
                product["reviews_3"] = q.get("3")
                product["reviews_2"] = q.get("2")
                product["reviews_1"] = q.get("1")
                for rev in payload.get("reviews") or []:
                    review_rows.append({
                        "sku": sku,
                        "review_id": rev.get("id"),
                        "date": rev.get("date"),
                        "user_name": rev.get("user_name"),
                        "rating": _to_float(rev.get("rating")),
                        "comment": rev.get("comment"),
                    })
        except Exception:
            pass

        with lock:
            append_jsonl(PRODUCTS_JSONL, product)
            for row in review_rows:
                append_jsonl(REVIEWS_JSONL, row)
            counters["seen"].add(sku)
            counters["processed"] += 1
            n = counters["processed"]
            if n % 25 == 0 or n <= 5:
                name = (product.get("name") or "")[:50]
                elapsed = time.time() - counters["start"]
                rate = n / elapsed if elapsed > 0 else 0
                print(f"[{idx}/{total}] OK {sku} ({n} done, {rate:.2f}/s) - {name}", flush=True)
    except Exception as e:
        print(f"[{idx}/{total}] EXC {type(e).__name__}: {e}", flush=True)
        with lock:
            counters["failed"] += 1


def crawl(limit: int, proxy: str | None = None, workers: int = 8):
    if proxy:
        print(f"[proxy] {proxy.split('@')[-1] if '@' in proxy else proxy}", flush=True)
    print(f"[sitemap] fetching {SITEMAP_URL}", flush=True)
    bootstrap = make_session(proxy=proxy)
    urls = fetch_sitemap_urls(bootstrap, limit=limit)
    print(f"[sitemap] {len(urls)} product URLs", flush=True)

    counters = {"processed": 0, "failed": 0,
                "seen": load_seen_skus(PRODUCTS_JSONL),
                "start": time.time()}
    lock = threading.Lock()
    total = len(urls)
    todo = []
    for i, url in enumerate(urls, 1):
        m = re.search(r"-(\d+)/?$", url)
        sku_guess = m.group(1) if m else None
        if sku_guess and sku_guess in counters["seen"]:
            continue
        todo.append((i, url))

    print(f"[plan] todo={len(todo)}, workers={workers}", flush=True)

    with ThreadPoolExecutor(max_workers=workers) as pool:
        futures = [pool.submit(_process_one, url, i, total, proxy, counters, lock) for i, url in todo]
        for _ in as_completed(futures):
            pass

    elapsed = time.time() - counters["start"]
    rate = counters["processed"] / elapsed if elapsed > 0 else 0
    print(f"\n[done] processed={counters['processed']}, failed={counters['failed']}, "
          f"elapsed={elapsed:.1f}s, rate={rate:.2f} SKU/s", flush=True)
    export_xlsx()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--limit", type=int, default=50, help="max products (0 = unlimited)")
    ap.add_argument("--proxy", type=str, default=None, help="proxy url, e.g. socks5h://user:pass@host:port")
    ap.add_argument("--workers", type=int, default=8, help="concurrent workers")
    ap.add_argument("--export-only", action="store_true", help="rebuild xlsx from jsonl and exit")
    args = ap.parse_args()

    if args.export_only:
        export_xlsx()
        return
    crawl(args.limit, proxy=args.proxy, workers=args.workers)


if __name__ == "__main__":
    main()
